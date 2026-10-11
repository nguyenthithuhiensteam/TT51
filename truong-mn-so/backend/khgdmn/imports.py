"""Đọc tệp Word/Excel để xem trước khi nhập — chuyển từ electron/import-word.cjs (mammoth) và import-excel.cjs (SheetJS)."""
from __future__ import annotations

import io
import re
from datetime import date, datetime, time

import mammoth
from openpyxl import load_workbook

from .validation import detect_personal_data


def decode_html(value) -> str:
    text = re.sub(r"<br\s*/?\s*>", "\n", str(value or ""), flags=re.IGNORECASE)
    text = text.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&#39;", "'")
    return re.sub(r"<[^>]+>", "", text).strip()


def extract_tables(html: str) -> list[list[list[str]]]:
    tables = []
    for table in re.findall(r"<table[^>]*>([\s\S]*?)</table>", str(html), flags=re.IGNORECASE):
        rows = []
        for row in re.findall(r"<tr[^>]*>([\s\S]*?)</tr>", table, flags=re.IGNORECASE):
            rows.append([decode_html(cell) for cell in re.findall(r"<t[dh][^>]*>([\s\S]*?)</t[dh]>", row, flags=re.IGNORECASE)])
        tables.append(rows)
    return tables


def parse_word_buffer(content: bytes, filename: str = "") -> dict:
    raw = mammoth.extract_raw_text(io.BytesIO(content))
    converted = mammoth.convert_to_html(io.BytesIO(content))
    html = converted.value or ""
    headings = [{"level": int(level), "text": decode_html(text)} for level, text in re.findall(r"<h([1-6])[^>]*>([\s\S]*?)</h\1>", html, flags=re.IGNORECASE)]
    paragraphs = [p for p in (decode_html(match) for match in re.findall(r"<p[^>]*>([\s\S]*?)</p>", html, flags=re.IGNORECASE)) if p]
    text = str(raw.value or "").strip()
    return {
        "filename": filename, "text": text, "headings": headings, "paragraphs": paragraphs, "tables": extract_tables(html),
        "warnings": [message.message for message in [*(raw.messages or []), *(converted.messages or [])]],
        "privacyWarnings": detect_personal_data(text),
    }


def _cell_text(value) -> str:
    """Gần với sheet_to_json(raw:false) của SheetJS: chuỗi đã định dạng, số nguyên không có '.0'."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y") if value.time() == time(0) else value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def _xlsx_sheets(content: bytes) -> list[tuple[str, list[list[str]]]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    return sheets


def _xls_sheets(content: bytes) -> list[tuple[str, list[list[str]]]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(_cell_text(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(_cell_text(bool(cell.value)))
                else:
                    row.append(_cell_text(cell.value))
            rows.append(row)
        sheets.append((sheet.name, rows))
    return sheets


def parse_excel_buffer(content: bytes, filename: str = "") -> dict:
    raw_sheets = _xls_sheets(content) if str(filename).lower().endswith(".xls") else _xlsx_sheets(content)
    sheets = []
    for name, rows in raw_sheets:
        # Bỏ hàng trống ở cuối (SheetJS dừng ở ô cuối có dữ liệu) và đệm các hàng cho cùng độ rộng.
        while rows and not any(rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, v in enumerate(row) if v), default=0) for row in rows), default=0)
        normalized = [(row + [""] * width)[:width] for row in rows]
        preview = normalized[:21]
        sheets.append({"name": name, "rowCount": len(normalized), "headers": preview[0] if preview else [], "preview": preview, "rows": normalized})
    joined = " ".join(cell for sheet in sheets for row in sheet["preview"] for cell in row)
    return {"filename": filename, "sheets": sheets, "privacyWarnings": detect_personal_data(joined)}
