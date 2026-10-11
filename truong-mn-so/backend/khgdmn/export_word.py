"""Xuất kế hoạch ra Word — chuyển từ electron/export-word.cjs (thư viện docx của Node) sang python-docx.

Giữ nguyên bố cục: khổ A4 (ngang cho năm/chủ đề/tuần, dọc cho giáo án), Times New Roman, bảng viền đơn,
hàng tiêu đề tô màu và lặp lại khi sang trang, khối chữ ký tối đa 3 người.
"""
from __future__ import annotations

import base64
import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor, Twips

from .plan_schema import DEVELOPMENT_DOMAINS, WEEK_DAYS, js_str, normalize_plan_type, parse_annual_rows, parse_theme_rows, parse_weekly_rows, split_lines

FONT = "Times New Roman"
A4_WIDTH, A4_HEIGHT = 11906, 16838
DOTS = "........................................................................................................................"
TYPE_TITLES = {
    "annual": "KẾ HOẠCH GIÁO DỤC NĂM",
    "theme": "KẾ HOẠCH GIÁO DỤC CHỦ ĐỀ",
    "weekly": "KẾ HOẠCH CHĂM SÓC GIÁO DỤC TRẺ",
    "lesson": "GIÁO ÁN HOẠT ĐỘNG GIÁO DỤC",
}


def _text(value) -> str:
    return js_str(value) if value not in (None, "", 0, False) else ""


def _set_run_font(run, size_half_points=24, bold=False, italics=False):
    run.font.name = FONT
    run.font.size = Pt(size_half_points / 2)
    run.font.bold = bold
    run.font.italic = italics
    run.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attribute), FONT)


def _fill_paragraph(paragraph, text="", *, size=24, bold=False, italics=False, center=False, right=False,
                    before=0, after=80, line=276, keep_next=False):
    for index, line_text in enumerate(re.split(r"\r?\n", _text(text))):
        run = paragraph.add_run()
        if index:
            run.add_break()
        run.add_text(line_text)
        _set_run_font(run, size, bold, italics)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT if right else WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.space_before = Twips(before)
    fmt.space_after = Twips(after)
    fmt.line_spacing = line / 240
    fmt.keep_with_next = keep_next
    return paragraph


def _shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _cell_margins(cell, top=90, right=100, bottom=90, left=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("right", right), ("bottom", bottom), ("left", left)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _insert_tbl_pr(tbl_pr, element):
    """Word kiểm tra thứ tự phần tử trong tblPr: tblBorders/tblLayout phải nằm trước tblCellMar/tblLook."""
    anchor = tbl_pr.find(qn("w:tblCellMar"))
    if anchor is None:
        anchor = tbl_pr.find(qn("w:tblLook"))
    if anchor is not None:
        anchor.addprevious(element)
    else:
        tbl_pr.append(element)


def _table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    _insert_tbl_pr(tbl_pr, borders)


def _row_flags(row, header=False, cant_split=True):
    tr_pr = row._tr.get_or_add_trPr()
    if cant_split:
        tr_pr.append(OxmlElement("w:cantSplit"))
    if header:
        tr_pr.append(OxmlElement("w:tblHeader"))


class _Builder:
    def __init__(self, landscape: bool):
        self.document = Document()
        self._styles()
        section = self.document.sections[0]
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = Twips(A4_HEIGHT), Twips(A4_WIDTH)
            margins = dict(top=650, right=650, bottom=650, left=650, header=300, footer=300)
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = Twips(A4_WIDTH), Twips(A4_HEIGHT)
            margins = dict(top=900, right=850, bottom=850, left=1100, header=350, footer=350)
        section.top_margin, section.right_margin = Twips(margins["top"]), Twips(margins["right"])
        section.bottom_margin, section.left_margin = Twips(margins["bottom"]), Twips(margins["left"])
        section.header_distance, section.footer_distance = Twips(margins["header"]), Twips(margins["footer"])

    def _styles(self):
        normal = self.document.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(12)
        normal.paragraph_format.line_spacing = 276 / 240
        normal.paragraph_format.space_after = Twips(60)

    def paragraph(self, text="", **options):
        return _fill_paragraph(self.document.add_paragraph(), text, **options)

    def table(self, widths, rows_spec):
        """rows_spec: danh sách hàng; mỗi hàng là dict {cells: [...], header, cant_split}.
        Mỗi ô: dict {text|paragraphs, bold, center, shading, span, valign}."""
        table = self.document.add_table(rows=0, cols=len(widths))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_w.set(qn("w:w"), str(sum(widths)))
        tbl_w.set(qn("w:type"), "dxa")
        _table_borders(table)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        _insert_tbl_pr(tbl_pr, layout)
        for grid_col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            grid_col.set(qn("w:w"), str(width))
        for spec in rows_spec:
            row = table.add_row()
            _row_flags(row, header=spec.get("header", False), cant_split=spec.get("cant_split", True))
            column = 0
            for cell_spec in spec["cells"]:
                span = cell_spec.get("span", 1)
                cell = row.cells[column]
                if span > 1:
                    cell = cell.merge(row.cells[column + span - 1])
                # Thứ tự trong tcPr: tcW, gridSpan, shd, tcMar, vAlign — shd/tcMar được nối vào trước vAlign.
                if cell_spec.get("shading"):
                    _shade(cell, cell_spec["shading"])
                _cell_margins(cell)
                cell.width = Twips(sum(widths[column:column + span]))
                cell.vertical_alignment = cell_spec.get("valign", WD_CELL_VERTICAL_ALIGNMENT.CENTER)
                first = cell.paragraphs[0]
                builders = cell_spec.get("paragraphs")
                if builders:
                    builders[0](first)
                    for build in builders[1:]:
                        build(cell.add_paragraph())
                else:
                    _fill_paragraph(first, cell_spec.get("text", ""), bold=cell_spec.get("bold", False), center=cell_spec.get("center", False))
                column += span
        return table


def _header(labels, widths):
    return {"header": True, "cells": [{"text": label, "bold": True, "center": True, "shading": "E6F2F1"} for label in labels]}


def _title_block(b: _Builder, plan: dict, school: dict, plan_type: str):
    b.paragraph(_text(school.get("name") or "TRƯỜNG MẦM NON").upper(), bold=True, center=True, size=26, after=20)
    b.paragraph(TYPE_TITLES[plan_type], bold=True, center=True, size=30, after=20)
    if plan_type == "weekly":
        b.paragraph(f"(TUẦN {_text(plan.get('weekNumber'))})" if plan.get("weekNumber") else "", bold=True, center=True, size=26, after=20)
    b.paragraph(_text(plan.get("title") or "Kế hoạch giáo dục"), bold=True, center=True, size=28, after=20)
    b.paragraph(f"Thời gian thực hiện: {_text(plan.get('period')) or '................................'}", bold=True, center=True, size=24, after=20)
    b.paragraph(f"Lớp: {_text(plan.get('className') or plan.get('ageGroup'))}    -    Giáo viên: {_text(plan.get('author'))}", italics=True, center=True, size=23, after=120)


def _annual(b: _Builder, plan: dict):
    widths = [520, 1120, 1800, 2400, 2400, 2440]
    rows = [_header(["STT", "Thời gian", "Chủ đề", "Mục tiêu", "Nội dung", "Hoạt động"], widths)]
    for index, item in enumerate(parse_annual_rows(plan)):
        rows.append({"cells": [{"text": index + 1, "center": True}, {"text": item["period"]}, {"text": item["theme"], "bold": True},
                               {"text": item["objective"]}, {"text": item["content"]}, {"text": item["activity"]}]})
    b.table(widths, rows)


def _theme(b: _Builder, plan: dict):
    widths = [560, 3360, 3240, 3240]
    rows = [_header(["STT", "Mục tiêu", "Nội dung", "Hoạt động"], widths)]
    items = parse_theme_rows(plan)
    sequence = 0
    for domain in DEVELOPMENT_DOMAINS:
        domain_rows = [item for item in items if item["domain"] == domain]
        if not domain_rows:
            continue
        rows.append({"cells": [{"text": domain.upper(), "bold": True, "shading": "F2F2F2", "span": 4}]})
        for item in domain_rows:
            sequence += 1
            rows.append({"cells": [{"text": sequence, "center": True}, {"text": item["objective"]}, {"text": item["content"]}, {"text": item["activity"]}]})
    b.table(widths, rows)
    if plan.get("themeFamilyWeeks"):
        b.paragraph("RÈN NỀN NẾP, THÓI QUEN VÀ PHỐI HỢP VỚI CHA MẸ TRẺ", bold=True, center=True, size=26, before=160, keep_next=True)
        family_rows = [_header(["Tuần", "Nội dung"], [1800, 9040])]
        for line in split_lines(plan["themeFamilyWeeks"]):
            week, *content = line.split("|")
            family_rows.append({"cant_split": False, "cells": [{"text": week, "bold": True}, {"text": "|".join(content).strip()}]})
        b.table([1800, 9040], family_rows)


def _weekly(b: _Builder, plan: dict):
    widths = [1600, 1848, 1848, 1848, 1848, 1848]
    rows = [_header(["Nội dung hoạt động", *(label for _, label in WEEK_DAYS)], widths)]
    for item in parse_weekly_rows(plan):
        rows.append({"cant_split": False, "cells": [{"text": item["activity"], "bold": True, "center": True}, *({"text": item[key]} for key, _ in WEEK_DAYS)]})
    if plan.get("objectives"):
        b.paragraph("MỤC TIÊU TRỌNG TÂM TRONG TUẦN", bold=True, size=25, keep_next=True)
        b.paragraph(plan["objectives"], after=120)
    morning = "\n".join(_text(plan.get(key)) for key in ("weeklyWelcome", "weeklyCircleTime", "weeklyMorningExercise") if plan.get(key))
    if morning:
        b.paragraph("ĐÓN TRẺ, TRÒ CHUYỆN, THỂ DỤC SÁNG", bold=True, size=25, keep_next=True)
        b.paragraph(morning, after=120)
    b.paragraph("HOẠT ĐỘNG HỌC", bold=True, size=25, keep_next=True, before=120 if morning else 0)
    b.table(widths, rows)
    for key, heading in (("weeklyCornerActivities", "HOẠT ĐỘNG GÓC"), ("weeklyOutdoorActivities", "HOẠT ĐỘNG NGOÀI TRỜI"),
                         ("weeklyMealSleep", "VỆ SINH, ĂN, NGỦ"), ("weeklyAfternoon", "HOẠT ĐỘNG CHIỀU"),
                         ("assessment", "ĐÁNH GIÁ VÀ ĐIỀU CHỈNH SAU TUẦN")):
        if plan.get(key):
            b.paragraph(heading, bold=True, size=25, before=160, keep_next=True)
            b.paragraph(plan[key])
    if plan.get("weeklyNotes"):
        b.paragraph("GHI CHÚ VÀ ĐIỀU CHỈNH TRONG TUẦN", bold=True, center=True, size=26, before=160, keep_next=True)
        b.paragraph(plan["weeklyNotes"])


def _lesson(b: _Builder, plan: dict):
    b.paragraph(plan.get("lessonDomain") or plan.get("developmentArea") or "LĨNH VỰC PHÁT TRIỂN", bold=True, center=True, size=27, after=20)
    b.paragraph(_text(plan.get("lessonTitle") or plan.get("title") or "HOẠT ĐỘNG GIÁO DỤC").upper(), bold=True, center=True, size=29, after=120)
    b.paragraph("I. MỤC ĐÍCH - YÊU CẦU", bold=True, size=26, keep_next=True)
    b.paragraph(plan.get("lessonObjectives") or plan.get("objectives") or "Chưa có nội dung.")
    b.paragraph("II. CHUẨN BỊ", bold=True, size=26, keep_next=True)
    b.paragraph(plan.get("lessonPreparation") or plan.get("materials") or "Chưa có nội dung.")
    b.paragraph("III. TIẾN HÀNH HOẠT ĐỘNG", bold=True, size=26, keep_next=True)
    widths = [5000, 5000]
    top = WD_CELL_VERTICAL_ALIGNMENT.TOP
    b.table(widths, [
        _header(["Hoạt động của cô", "Hoạt động của trẻ"], widths),
        {"cant_split": False, "cells": [{"text": plan.get("teacherActivities") or plan.get("activities") or "Chưa có nội dung.", "valign": top},
                                        {"text": plan.get("childActivities") or "", "valign": top}]},
    ])
    b.paragraph("ĐÁNH GIÁ CUỐI NGÀY", bold=True, center=True, size=26, before=180, keep_next=True)
    b.paragraph(plan.get("dailyEvaluation") or plan.get("assessment") or f"{DOTS}\n{DOTS}\n{DOTS}")


def _signature_image(paragraph, signer: dict):
    match = re.match(r"^data:image/(png|jpeg);base64,(.+)$", _text(signer.get("imageData")), re.DOTALL)
    if signer.get("mode") != "image" or not match:
        _fill_paragraph(paragraph, "\n\n", center=True)
        return
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def px(value, fallback):
        try:
            return int(float(value)) or fallback
        except (TypeError, ValueError):
            return fallback

    run = paragraph.add_run()
    run.add_picture(io.BytesIO(base64.b64decode(match.group(2))), width=Emu(px(signer.get("width"), 120) * 9525), height=Emu(px(signer.get("height"), 55) * 9525))


def _signatures(b: _Builder, signers: list):
    if not signers:
        b.paragraph("NHẬN XÉT CỦA BAN GIÁM HIỆU", bold=True, center=True, size=26, before=180)
        b.paragraph(f"{DOTS}\n{DOTS}\n{DOTS}")
        return
    shown = signers[:3]
    widths = [10000 // len(shown)] * len(shown)
    cells = []
    for signer in shown:
        cells.append({"paragraphs": [
            lambda p, s=signer: _fill_paragraph(p, s.get("displayTitle") or s.get("role"), bold=True, center=True),
            lambda p, s=signer: _signature_image(p, s),
            lambda p, s=signer: _fill_paragraph(p, "" if s.get("mode") == "hand" else s.get("name"), bold=True, center=True),
        ]})
    b.paragraph("XÁC NHẬN VÀ CHỮ KÝ", bold=True, center=True, size=26, before=180)
    b.table(widths, [{"cells": cells}])


def create_plan_docx(plan: dict | None = None, school: dict | None = None) -> bytes:
    plan = plan or {}
    school = school or {}
    plan_type = normalize_plan_type(plan.get("level"))
    builder = _Builder(landscape=plan_type != "lesson")
    _title_block(builder, plan, school, plan_type)
    {"annual": _annual, "theme": _theme, "weekly": _weekly, "lesson": _lesson}[plan_type](builder, plan)
    _signatures(builder, plan.get("signers") or [])
    output = io.BytesIO()
    builder.document.save(output)
    return output.getvalue()


def create_review_docx(report: dict | None = None, school: dict | None = None) -> bytes:
    report = report or {}
    return create_plan_docx({
        "level": "Ngày/hoạt động", "title": report.get("title") or "Báo cáo rà soát dữ liệu", "period": report.get("period"),
        "author": report.get("author"), "lessonDomain": "BÁO CÁO RÀ SOÁT", "lessonTitle": report.get("title") or "Báo cáo rà soát dữ liệu",
        "lessonObjectives": report.get("summary"), "lessonPreparation": report.get("progress"), "teacherActivities": report.get("issues"),
        "childActivities": "", "dailyEvaluation": report.get("conclusion"),
    }, school)
