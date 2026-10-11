"""Xuất Excel — chuyển từ electron/export-excel.cjs (exceljs) sang openpyxl."""
from __future__ import annotations

import io
import json

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .plan_schema import js_str

SHEETS = (("Kho dữ liệu mở", "records"), ("Ma trận mục tiêu", "matrix"), ("Danh sách kế hoạch", "plans"),
          ("Tiến độ rà soát", "reviewProgress"), ("Cảnh báo dữ liệu", "issues"))
HEADER_FILL = PatternFill("solid", fgColor="FF168C7D")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")


def _cell_value(value):
    """Giữ nguyên số/bool; mảng/object chuyển thành chuỗi; bỏ ký tự điều khiển mà openpyxl không ghi được."""
    if value is None:
        return ""
    if isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, list):
        value = ", ".join(js_str(item) for item in value)
    elif isinstance(value, dict):
        value = json.dumps(value, ensure_ascii=False)
    return ILLEGAL_CHARACTERS_RE.sub("", js_str(value))


def _values_for(rows):
    if not isinstance(rows, list) or not rows:
        return ["Thông tin"], [["Chưa có dữ liệu"]]
    headers = list(dict.fromkeys(key for row in rows for key in (row or {}).keys()))
    return headers, [[_cell_value((row or {}).get(header)) for header in headers] for row in rows]


def _style_header(sheet, count: int, center: bool = True):
    for column in range(1, count + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        if center:
            cell.alignment = Alignment(vertical="center", horizontal="center")


def _save(workbook: Workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_professional_workbook(data: dict | None = None) -> bytes:
    data = data or {}
    school = data.get("school") or {}
    school_name = school.get("name") or "Trường Mầm non Số"
    workbook = Workbook()
    workbook.properties.creator = school_name
    guide = workbook.active
    guide.title = "Hướng dẫn"
    for line in (f"{school_name} — NĂM HỌC {school.get('schoolYear') or '2026–2027'}",
                 f"Đơn vị chủ quản: {school.get('governingBody') or 'Chưa cấu hình'}",
                 "Mỗi sheet là một nhóm dữ liệu chuyên môn.",
                 "Cột ID, nguồn PDF và trang nguồn dùng để truy xuất ngược.",
                 "PDF chỉ là tài liệu nguồn, không phải dữ liệu làm việc."):
        guide.append([line])
    guide.column_dimensions["A"].width = 90
    guide["A1"].font = Font(bold=True, size=14)
    for title, key in SHEETS:
        sheet = workbook.create_sheet(title)
        headers, rows = _values_for(data.get(key))
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        _style_header(sheet, len(headers))
        for index, header in enumerate(headers):
            longest = max([len(js_str(header)) or 10, *(len(js_str(row[index])) for row in rows[:100])])
            sheet.column_dimensions[get_column_letter(index + 1)].width = min(45, max(12, longest + 2))
    return _save(workbook)


def _join(values) -> str:
    return "; ".join(js_str(v) for v in (values or []))


def create_directory_workbook(kind: str, rows=None, template=False, school: dict | None = None) -> bytes:
    rows = rows or []
    school = school or {}
    is_class = kind == "classes"
    is_video = kind == "videos"
    title = "Danh mục lớp" if is_class else "Danh mục video" if is_video else "Danh mục giáo viên"
    if is_class:
        headers = ["ID lớp", "Tên lớp", "Nhóm độ tuổi", "Năm học", "Điểm trường", "Số lượng trẻ", "ID giáo viên phụ trách", "ID giáo viên phối hợp", "Trạng thái", "Ghi chú"]
    elif is_video:
        headers = ["ID", "Tiêu đề", "Mô tả", "Nhóm nội dung", "Đối tượng", "Loại nguồn", "URL HTTPS", "Thời lượng", "Thứ tự", "Trạng thái", "Khóa ngữ cảnh"]
    else:
        headers = ["ID", "Họ và tên", "Chức vụ", "Vai trò", "ID lớp phụ trách", "Tổ chuyên môn", "Trạng thái", "Liên hệ"]
    workbook = Workbook()
    workbook.properties.creator = school.get("name") or "Trường Mầm non Số"
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    if template:
        if is_class:
            sheet.append(["lop-01", "Lớp mẫu", "4–5 tuổi", school.get("schoolYear") or "2026–2027", "", 25, "gv-01", "", "Đang sử dụng", ""])
        elif is_video:
            sheet.append(["", "Nhập dữ liệu từ Excel", "Hướng dẫn thao tác", "Quản lý dữ liệu", "teacher; team_lead", "Trực tuyến", "https://example.com/video", "05:00", 1, "Bật", "open-data"])
        else:
            sheet.append(["gv-01", "Nguyễn Văn A", "Giáo viên", "Người soạn; Giáo viên phối hợp", "lop-01", "Tổ Mẫu giáo", "Đang công tác", ""])
    else:
        for row in rows:
            g = row.get
            if is_class:
                values = [g("id"), g("name"), g("ageGroup"), g("schoolYear"), g("campus"), g("childCount"), g("leadTeacherId"),
                          _join(g("collaboratingTeacherIds")), "Đang sử dụng" if g("active") else "Ngừng sử dụng", g("notes")]
            elif is_video:
                values = [g("id"), g("title"), g("description"), g("category"), _join(g("audience")),
                          "Ngoại tuyến" if g("source_type") == "offline" else "Trực tuyến", g("url") or "", g("duration"), g("sort_order"),
                          "Bật" if g("enabled") else "Tắt", g("context_key") or ""]
            else:
                values = [g("id"), g("name"), g("title"), _join(g("roles")), _join(g("classIds")), g("team"),
                          "Đang công tác" if g("active") else "Ngừng sử dụng", g("contact")]
            sheet.append([_cell_value(v) for v in values])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    _style_header(sheet, len(headers), center=False)
    for index, header in enumerate(headers):
        sheet.column_dimensions[get_column_letter(index + 1)].width = min(38, max(14, len(header) + 3))
    guide = workbook.create_sheet("Hướng dẫn")
    for line in (f"{title.upper()} – {school.get('name') or 'Trường Mầm non Số'}",
                 "Không đổi tên hàng tiêu đề. ID phải duy nhất. Nhiều vai trò hoặc ID được phân cách bằng dấu chấm phẩy.",
                 "Dữ liệu được nhập và lưu cục bộ trên máy."):
        guide.append([line])
    guide.column_dimensions["A"].width = 100
    return _save(workbook)
