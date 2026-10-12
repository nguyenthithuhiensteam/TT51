"""Kiểm thử xuất Word/Excel/PDF — đối chiếu nội dung Word với bản Node khi có Node."""
import io
import json
import os
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from backend.khgdmn.export_word import create_plan_docx
from test_khgdmn_core import _docx_bytes, client, login, setup_admin  # noqa: F401  (fixture dùng chung)


def setup_principal(client):
    """Quản trị và một giáo viên (vai trò khác như hiệu trưởng nay quy về giáo viên)."""
    admin = setup_admin(client)
    client.post("/api/accounts", headers=admin, json={"username": "hieutruong", "fullName": "Hiệu trưởng", "password": "matkhau123",
                                                      "roles": ["principal"], "mustChangePassword": False})
    return admin, login(client, "hieutruong", "matkhau123")

FIXTURES = Path(__file__).with_name("fixtures")
DATA = json.loads((FIXTURES / "plans.json").read_text(encoding="utf-8"))


def docx_text(content: bytes) -> list[str]:
    """Toàn bộ chữ theo thứ tự đoạn văn và ô bảng (bỏ dòng trống) để so sánh hai bản xuất."""
    document = Document(io.BytesIO(content))
    lines = []
    body = document.element.body
    for paragraph in body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
        text = "".join(node.text or "" for node in paragraph.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
        if text.strip():
            lines.append(text.strip())
    return lines


@pytest.mark.parametrize("plan", DATA["plans"], ids=[p["id"] for p in DATA["plans"]])
def test_word_export_opens_and_has_content(plan):
    content = create_plan_docx(plan, DATA["school"])
    document = Document(io.BytesIO(content))
    section = document.sections[0]
    assert (section.page_width > section.page_height) == (plan["level"] != "lesson")
    text = "\n".join(docx_text(content))
    assert "TRƯỜNG MẦM NON HƯỚNG DƯƠNG" in text


@pytest.mark.skipif(shutil.which("node") is None, reason="Không có Node để đối chiếu")
def test_word_export_matches_node_text(tmp_path):
    subprocess.run(["node", str(FIXTURES / "node_export.cjs"), str(tmp_path)], check=True, cwd=FIXTURES)
    for plan in DATA["plans"]:
        node_text = docx_text((tmp_path / f"{plan['id']}.docx").read_bytes())
        python_text = docx_text(create_plan_docx(plan, DATA["school"]))
        assert python_text == node_text, plan["id"]


def test_export_routes(client):
    admin, principal = setup_principal(client)
    plan = DATA["plans"][0]
    assert client.post("/api/exports/word", headers=admin, json={"filename": "k.docx", "payload": plan}).status_code == 200
    word = client.post("/api/exports/word", headers=principal, json={"filename": "Kế hoạch năm.docx", "payload": plan, "school": DATA["school"]})
    assert word.status_code == 200, word.text
    assert word.headers["content-disposition"] == "attachment; filename*=UTF-8''K%E1%BA%BF%20ho%E1%BA%A1ch%20n%C4%83m.docx"
    Document(io.BytesIO(word.content))
    bad_name = client.post("/api/exports/word", headers=principal, json={"filename": "a.pdf", "payload": plan})
    assert bad_name.json()["error"] == "Tên tệp phải có phần mở rộng .docx."
    review = client.post("/api/exports/word", headers=principal, json={"filename": "r.docx", "type": "review", "payload": {"title": "Rà soát", "summary": "Tóm tắt"}})
    assert "BÁO CÁO RÀ SOÁT" in "\n".join(docx_text(review.content))

    excel = client.post("/api/exports/excel", headers=principal, json={"filename": "kho.xlsx", "payload": {"school": DATA["school"], "plans": [{"id": "p1", "title": "KH", "tags": ["a", "b"]}]}})
    workbook = load_workbook(io.BytesIO(excel.content))
    assert workbook.sheetnames == ["Hướng dẫn", "Kho dữ liệu mở", "Ma trận mục tiêu", "Danh sách kế hoạch", "Tiến độ rà soát", "Cảnh báo dữ liệu"]
    assert [c.value for c in workbook["Danh sách kế hoạch"][2]] == ["p1", "KH", "a, b"]
    assert workbook["Kho dữ liệu mở"]["A2"].value == "Chưa có dữ liệu"

    directory = client.post("/api/exports/directory", headers=admin, json={"filename": "lop.xlsx", "kind": "classes", "template": True, "school": DATA["school"]})
    assert load_workbook(io.BytesIO(directory.content))["Danh mục lớp"]["B2"].value == "Lớp mẫu"
    assert client.post("/api/exports/directory", headers=admin, json={"filename": "x.xlsx", "kind": "khac"}).json()["error"] == "Loại danh mục không hợp lệ."


def test_custom_template_used_for_word_export(client):
    admin, principal = setup_principal(client)
    imported = client.post("/api/templates/import", headers=admin, files={"file": ("mau.docx", _docx_bytes("{{PLAN_TABLE}}"))}, data={"type": "annual"}).json()
    template_id = next(t["id"] for t in imported["templates"] if not t["builtIn"])
    plan = {**DATA["plans"][0], "templateId": template_id}
    word = client.post("/api/exports/word", headers=principal, json={"filename": "k.docx", "payload": plan, "school": DATA["school"]})
    text = docx_text(word.content)
    assert text[:6] == ["STT", "Thời gian", "Chủ đề", "Mục tiêu", "Nội dung", "Hoạt động"]
    assert "Trường mầm non" in text


@pytest.mark.skipif(not (os.environ.get("PLAYWRIGHT_CHANNEL") or os.environ.get("PLAYWRIGHT_CHROMIUM_PATH")), reason="Chưa cấu hình trình duyệt in PDF")
def test_pdf_export(client):
    _admin, principal = setup_principal(client)
    response = client.post("/api/exports/pdf", headers=principal, json={"filename": "giao-an.pdf", "payload": DATA["plans"][3], "school": DATA["school"]})
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"%PDF")
