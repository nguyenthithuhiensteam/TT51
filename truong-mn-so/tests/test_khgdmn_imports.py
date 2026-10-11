"""Kiểm thử nhập Word/Excel/PDF và kho tài liệu — đối chiếu kết quả đọc với bản Node khi có Node."""
import io
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from backend.khgdmn.imports import parse_excel_buffer, parse_word_buffer
from test_khgdmn_core import client, setup_admin  # noqa: F401

FIXTURES = Path(__file__).with_name("fixtures")
MINIMAL_PDF = b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n"


def sample_docx() -> bytes:
    document = Document()
    document.add_heading("Kế hoạch tuần 1", level=1)
    document.add_paragraph("Mục tiêu: trẻ biết rửa tay. Liên hệ 0912 345 678.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Thứ", "Hoạt động"
    table.cell(1, 0).text, table.cell(1, 1).text = "Thứ hai", "Kể chuyện"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def sample_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sách"
    sheet.append(["username", "fullName", "roles"])
    sheet.append(["gv.lan", "Cô Lan", "teacher"])
    sheet.append(["gv.mai", "Cô Mai", 12])
    sheet.append([])
    sheet.append(["", "", "ghi chú cuối"])
    workbook.create_sheet("Trống")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def node_parse(kind: str, content: bytes, tmp_path: Path, name: str) -> dict:
    file = tmp_path / name
    file.write_bytes(content)
    output = subprocess.run(["node", str(FIXTURES / "node_parse.cjs"), kind, str(file)], check=True, capture_output=True, cwd=FIXTURES).stdout
    return json.loads(output.decode("utf-8"))


def test_word_preview():
    result = parse_word_buffer(sample_docx(), "kh.docx")
    assert result["headings"] == [{"level": 1, "text": "Kế hoạch tuần 1"}]
    assert result["tables"] == [[["Thứ", "Hoạt động"], ["Thứ hai", "Kể chuyện"]]]
    assert result["privacyWarnings"] == ["Có thể chứa số điện thoại."]


def test_excel_preview():
    result = parse_excel_buffer(sample_xlsx(), "ds.xlsx")
    first = result["sheets"][0]
    assert first["headers"] == ["username", "fullName", "roles"]
    assert first["rows"][2] == ["gv.mai", "Cô Mai", "12"]
    assert first["rowCount"] == 5 and first["rows"][3] == ["", "", ""]
    assert result["sheets"][1] == {"name": "Trống", "rowCount": 0, "headers": [], "preview": [], "rows": []}


@pytest.mark.skipif(shutil.which("node") is None, reason="Không có Node để đối chiếu")
def test_word_and_excel_match_node(tmp_path):
    for key in ("headings", "paragraphs", "tables", "text", "privacyWarnings"):
        assert parse_word_buffer(sample_docx(), "kh.docx")[key] == node_parse("word", sample_docx(), tmp_path, "kh.docx")[key], key
    python_excel = parse_excel_buffer(sample_xlsx(), "ds.xlsx")
    node_excel = node_parse("excel", sample_xlsx(), tmp_path, "ds.xlsx")
    assert python_excel["sheets"][0] == node_excel["sheets"][0]
    assert python_excel["privacyWarnings"] == node_excel["privacyWarnings"]


def test_import_routes_and_documents(client):
    admin = setup_admin(client)
    word = client.post("/api/imports/choose-files", headers=admin, data={"kind": "word"}, files=[("files", ("kh.docx", sample_docx()))]).json()
    assert word["canceled"] is False and word["files"][0]["preview"]["headings"][0]["text"] == "Kế hoạch tuần 1"
    wrong = client.post("/api/imports/choose-files", headers=admin, data={"kind": "word"}, files=[("files", ("a.pdf", MINIMAL_PDF))])
    assert wrong.json()["error"] == "Định dạng .pdf không được phép cho loại nhập này."

    chosen = client.post("/api/imports/choose-files", headers=admin, data={"kind": "pdf"}, files=[("files", ("Nguồn.pdf", MINIMAL_PDF)), ("files", ("Nguồn.pdf", MINIMAL_PDF))]).json()
    tokens = [f["token"] for f in chosen["files"]]
    imported = client.post("/api/imports/commit-pdf", headers=admin, json={"tokens": tokens, "actor": "Quản trị"}).json()
    assert [doc["relativePath"] for doc in imported] == ["user-added/Nguồn.pdf", "user-added/Nguồn-2.pdf"]
    again = client.post("/api/imports/commit-pdf", headers=admin, json={"tokens": tokens})
    assert again.json()["error"] == "Không còn tệp PDF hợp lệ để thêm."

    served = client.get("/api/documents/file", headers=admin, params={"relativePath": "user-added/Nguồn.pdf"})
    assert served.status_code == 200 and served.content == MINIMAL_PDF
    blocked = client.get("/api/documents/file", headers=admin, params={"relativePath": "../../README.md"})
    assert blocked.json()["error"] == "Ứng dụng đã chặn đường dẫn nằm ngoài kho tài liệu."
    program = client.get("/api/documents/file", headers=admin, params={"relativePath": "doc-010aa03b63c4.pdf"})
    assert program.status_code == 200 and program.content.startswith(b"%PDF")

    assert client.post("/api/imports/undo-pdf", headers=admin, json={"relativePaths": ["doc-010aa03b63c4.pdf"]}).json()["error"] == "Chỉ có thể hoàn tác PDF nguồn bổ sung."
    undone = client.post("/api/imports/undo-pdf", headers=admin, json={"relativePaths": ["user-added/Nguồn.pdf"]}).json()
    assert undone == {"ok": True, "removed": ["user-added/Nguồn.pdf"]}

    guide = client.get("/api/guides/pdf", headers=admin)
    assert guide.status_code == 200 and guide.content.startswith(b"%PDF")
    assert client.get("/api/guides/exe", headers=admin).json()["error"] == "Định dạng tài liệu hướng dẫn không hợp lệ."
    assert client.get("/api/resources/trusted/preschool-teacher-assistant", headers=admin).json()["url"].startswith("https://chatgpt.com/")
